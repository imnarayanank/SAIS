import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\Admin\Desktop\LEARNABLE\backend\uploads\timetable_1.xlsx", data_only=True)
for sheet_name in wb.sheetnames:
    print(f"--- Sheet: {sheet_name} ---")
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    for i, row in enumerate(rows[:20]):
        print(f"Row {i}: {row}")
