import openpyxl
import os
import sys

# Define standard VTOP times
VIT_THEORY_STARTS = ["08:00", "08:55", "09:50", "10:45", "11:40", "12:35", "Lunch", "14:00", "14:55", "15:50", "16:45", "17:40", "18:35"]
VIT_THEORY_ENDS   = ["08:50", "09:45", "10:40", "11:35", "12:30", "13:25", "Lunch", "14:50", "15:45", "16:40", "17:35", "18:30", "19:25"]
VIT_LAB_STARTS    = ["08:00", "08:50", "09:50", "10:40", "11:40", "12:30", "Lunch", "14:00", "14:50", "15:50", "16:40", "17:40", "18:30"]
VIT_LAB_ENDS      = ["08:50", "09:40", "10:40", "11:30", "12:30", "13:20", "Lunch", "14:50", "15:40", "16:40", "17:30", "18:30", "19:20"]

def parse_format_c_test(wb):
    course_map = {}
    
    # 1. Parse Course Details (if available)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        
        in_course_table = False
        course_col = -1
        faculty_col = -1
        
        for row in rows:
            if not in_course_table:
                row_strs = [str(c).strip().lower() for c in row if c is not None]
                if "course" in row_strs and "faculty details" in row_strs:
                    in_course_table = True
                    course_col = row_strs.index("course")
                    faculty_col = row_strs.index("faculty details")
            else:
                if not any(row):
                    continue
                if course_col < len(row) and faculty_col < len(row):
                    course_cell = row[course_col]
                    faculty_cell = row[faculty_col]
                    
                    if course_cell and faculty_cell and " - " in str(course_cell):
                        parts = str(course_cell).split(" - ", 1)
                        code = parts[0].strip()
                        name = parts[1].strip() if len(parts) > 1 else code
                        
                        fac_parts = str(faculty_cell).split(" - ", 1)
                        teacher = fac_parts[0].strip()
                        
                        course_map[code] = {"name": name, "teacher": teacher}

    # 2. Parse Grid
    entries = []
    days = ["MON", "TUE", "WED", "THU", "FRI", "SAT", "SUN"]
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        
        theory_starts = VIT_THEORY_STARTS
        theory_ends = VIT_THEORY_ENDS
        lab_starts = VIT_LAB_STARTS
        lab_ends = VIT_LAB_ENDS
        
        grid_start_row = -1
        day_col_idx = -1
        type_col_idx = -1
        data_col_idx = -1
        
        # Detect grid
        for i, row in enumerate(rows):
            # Check all cells to find "MON" followed by "THEORY"
            for j in range(len(row) - 1):
                if str(row[j]).strip().upper() in days and str(row[j+1]).strip().upper() == "THEORY":
                    grid_start_row = i
                    day_col_idx = j
                    type_col_idx = j + 1
                    data_col_idx = j + 2
                    break
            if grid_start_row != -1:
                break
                
        if grid_start_row != -1:
            current_day = None
            
            for i in range(grid_start_row, len(rows)):
                row = rows[i]
                if not any(row):
                    continue
                    
                cell_day = str(row[day_col_idx]).strip().upper() if day_col_idx < len(row) and row[day_col_idx] is not None else ""
                cell_type = str(row[type_col_idx]).strip().upper() if type_col_idx < len(row) and row[type_col_idx] is not None else ""
                
                if cell_day in days:
                    current_day = cell_day
                    
                if cell_type not in ["THEORY", "LAB"]:
                    continue
                    
                if not current_day:
                    continue
                    
                starts = theory_starts if cell_type == "THEORY" else lab_starts
                ends = theory_ends if cell_type == "THEORY" else lab_ends
                
                for col_idx in range(data_col_idx, len(row)):
                    cell = row[col_idx]
                    if not cell or str(cell).strip() == "" or str(cell).strip() == "-":
                        continue
                        
                    cell_text = str(cell).strip()
                    
                    if "-" in cell_text and len(cell_text.split("-")) > 2:
                        parts = cell_text.split("-")
                        slot = parts[0]
                        course_code = parts[1]
                        venue = parts[3] + "-" + parts[4] if len(parts) > 4 else (parts[3] if len(parts) > 3 else "")
                            
                        mapped = course_map.get(course_code, {})
                        subject_name = mapped.get("name", course_code)
                        teacher = mapped.get("teacher", "Unknown")
                        
                        time_idx = col_idx - data_col_idx
                        start_time = starts[time_idx] if time_idx < len(starts) else "00:00"
                        end_time = ends[time_idx] if time_idx < len(ends) else "00:00"
                        
                        if start_time.lower() != "lunch":
                            day_full = current_day.capitalize() if current_day != "MON" else "Monday"
                            if current_day == "TUE": day_full = "Tuesday"
                            if current_day == "WED": day_full = "Wednesday"
                            if current_day == "THU": day_full = "Thursday"
                            
                            entries.append({
                                "subject_name": subject_name,
                                "course_code": course_code,
                                "teacher": teacher,
                                "classroom": venue,
                                "day": day_full,
                                "start_time": start_time,
                                "end_time": end_time
                            })
                            
            if entries:
                return entries
    return []

wb = openpyxl.load_workbook(r"C:\Users\Admin\Desktop\LEARNABLE\backend\uploads\timetable_1.xlsx", data_only=True)
res = parse_format_c_test(wb)
print(f"Found {len(res)} entries")
for r in res[:5]:
    print(r)
